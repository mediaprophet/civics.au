"""Bounded importer for ABS SEIFA 2021 Table 1 summary workbooks.

Usage: python scripts/import-seifa-table1.py <workbook.xlsx> <lga|sa2> <output.json>
The original workbook stays outside the repository; this creates a governance-required
release package with an input hash and source-cell locators for review.
"""
import datetime
import hashlib
import json
import math
import sys
from pathlib import Path

import openpyxl

SOURCE_URL = "https://www.abs.gov.au/statistics/people/people-and-communities/socio-economic-indexes-areas-seifa-australia/2021"
MEASURES = [
    ("seifa-irsd-score", 3, "Index of Relative Socio-economic Disadvantage score"),
    ("seifa-irsad-score", 5, "Index of Relative Socio-economic Advantage and Disadvantage score"),
    ("seifa-ier-score", 7, "Index of Economic Resources score"),
    ("seifa-ieo-score", 9, "Index of Education and Occupation score"),
    ("usual-resident-population", 11, "Usual Resident Population"),
]


def fail(message):
    raise SystemExit(f"SEIFA import: {message}")


def main():
    if len(sys.argv) != 4:
        fail("usage: import-seifa-table1.py workbook.xlsx <lga|sa2> output.json")
    source, geography, destination = Path(sys.argv[1]), sys.argv[2], Path(sys.argv[3])
    if geography not in {"lga", "sa2"}:
        fail("geography must be lga or sa2")
    if not source.is_file():
        fail(f"workbook does not exist: {source}")
    workbook_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    observed_at = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    if "Table 1" not in workbook.sheetnames:
        fail("Table 1 sheet is required")
    sheet = workbook["Table 1"]
    if sheet.cell(6, 1).value is None or sheet.cell(6, 2).value is None:
        fail("expected Table 1 headers on row 6")
    code_width = 5 if geography == "lga" else 9
    release_id = f"abs_seifa_2021_{geography}_table_1"
    release = {
        "id": release_id,
        "datasetId": "D04-ABS-SEIFA",
        "publisher": "Australian Bureau of Statistics",
        "sourceUrl": SOURCE_URL,
        "retrievedAt": observed_at,
        "referencePeriod": "2021",
        "licenceStatus": "Confirm ABS conditions of use before public publication.",
        "sha256": workbook_hash,
        "schemaVersion": "ABS-SEIFA-2021-Table-1",
        "mappingVersion": "web-civics-seifa-table1-0.1.0",
        "access": "public",
        "admission": "governance-required",
        "localSourceObservedAt": observed_at,
        "originalAcquisitionEvidence": "Not supplied with the local file; confirm source download and licence before changing admission to staged.",
    }
    limitation = "SEIFA is an area-level context measure. It must not be used to infer an individual's circumstances, service demand, eligibility or priority."
    destination.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with destination.open("w", encoding="utf-8", newline="\n") as stream:
        stream.write('{"release":')
        json.dump(release, stream, separators=(",", ":"), ensure_ascii=False)
        stream.write(',"observations":[\n')
        first = True
        for row_number, row in enumerate(sheet.iter_rows(min_row=7, values_only=True), start=7):
            code, name = row[0], row[1]
            if code is None or not isinstance(code, (int, float)) or not name:
                continue
            code = f"{int(code):0{code_width}d}"
            geography_id = f"urn:abs:seifa-2021:{geography}:{code}"
            for indicator, column, _label in MEASURES:
                value = row[column - 1] if len(row) >= column else None
                if value is not None and (not isinstance(value, (int, float)) or not math.isfinite(value)):
                    value = None
                observation = {
                    "id": f"seifa_2021_{geography}_{code}_{indicator.replace('-', '_')}",
                    "indicatorId": indicator,
                    "originalValue": "not published" if value is None else str(value),
                    "unit": "persons" if indicator == "usual-resident-population" else "SEIFA index score",
                    "geographyId": geography_id,
                    "geographyVersion": f"2021 {'Local Government Area' if geography == 'lga' else 'Statistical Area Level 2'}",
                    "referencePeriod": "2021",
                    "sourceLocator": f"Table 1!{openpyxl.utils.get_column_letter(column)}{row_number}",
                    "evidenceState": "direct-count" if indicator == "usual-resident-population" else "context-only",
                    "limitations": [limitation, f"Area name in source: {name}."],
                }
                if value is not None:
                    observation["numericValue"] = value
                if not first:
                    stream.write(',\n')
                json.dump(observation, stream, separators=(",", ":"), ensure_ascii=False)
                first = False
                count += 1
        stream.write('\n]}\n')
    print(json.dumps({"releaseId": release_id, "observations": count, "workbookSha256": workbook_hash, "admission": release["admission"], "output": str(destination)}, indent=2))


if __name__ == "__main__":
    main()
